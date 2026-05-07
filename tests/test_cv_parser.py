import io

import pytest
from docx import Document

from app.pipeline.cv_parser import _normalize, extract_text


def make_docx(text: str) -> bytes:
    """Создаёт минимальный DOCX файл с заданным текстом."""
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestExtractText:
    def test_docx_extracts_text(self):
        content = "John Doe | Python Developer | 5 years experience"
        file_bytes = make_docx(content)
        result = extract_text(file_bytes, "cv.docx")
        assert "John Doe" in result
        assert "Python Developer" in result

    def test_unsupported_format_raises(self):
        with pytest.raises(ValueError, match="Unsupported file format"):
            extract_text(b"fake content", "photo.jpg")

    def test_unsupported_format_png_raises(self):
        with pytest.raises(ValueError, match="Unsupported file format"):
            extract_text(b"fake content", "resume.png")

    def test_empty_pdf_raises(self):
        # Невалидный PDF → pdfplumber бросает исключение, которое мы ловим как ValueError
        fake_pdf = b"not a real pdf content at all"
        with pytest.raises((ValueError, Exception)):
            extract_text(fake_pdf, "scan.pdf")

    def test_docx_multiline(self):
        doc = Document()
        doc.add_paragraph("Skills: Python, FastAPI")
        doc.add_paragraph("Experience: 3 years at Acme Corp")
        buf = io.BytesIO()
        doc.save(buf)
        result = extract_text(buf.getvalue(), "cv.docx")
        assert "Python" in result
        assert "Acme Corp" in result


class TestNormalize:
    def test_collapses_multiple_spaces(self):
        result = _normalize("hello   world")
        assert result == "hello world"

    def test_collapses_multiple_newlines(self):
        result = _normalize("line1\n\n\n\nline2")
        assert result == "line1\n\nline2"

    def test_strips_leading_trailing(self):
        result = _normalize("  hello  ")
        assert result == "hello"

    def test_preserves_double_newline(self):
        result = _normalize("section1\n\nsection2")
        assert result == "section1\n\nsection2"
